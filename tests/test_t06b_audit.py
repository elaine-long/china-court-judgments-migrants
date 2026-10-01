"""T06b §5: audit workbook structure, validation and blindness."""
import re

import pytest
from openpyxl import load_workbook

from src import paths

ANN = paths.ROOT / "annotation"
BOOK = ANN / "T05_audit.xlsx"
pytestmark = pytest.mark.skipif(not BOOK.exists(), reason="audit workbook not built")

CORE = ["origin_text", "nonlocal_origin", "residence_local", "status_at_judgment", "ever_bail", "plea_formal",
        "jiejie", "procedure", "counsel_any", "penalty_type", "term_months", "probation"]
FREE = {"origin_text", "term_months"}


def test_sheets_and_columns():
    wb = load_workbook(BOOK)
    assert {"disputes", "audit", "codebook", "log"} <= set(wb.sheetnames)
    for name in ("disputes", "audit"):
        head = [c.value for c in wb[name][1]]
        assert head[:3] == ["uid", "court_name", "judgment_date"]
        expected = [f"{f}_{s}" for f in CORE for s in ("prefill", "final", "evidence")]
        assert head[3:3 + 36] == expected
        assert head[-4:] == ["changed", "flag", "note", "minutes"]
        assert not any(re.search(r"source|rule|llm|model|stratum|doc_id|weight", str(h)) for h in head)


def test_counts_and_final_equals_prefill():
    wb = load_workbook(BOOK)
    audit = list(wb["audit"].iter_rows(min_row=2, values_only=True))
    assert len(audit) == 150 and [r[0] for r in audit][:2][0].startswith("u")
    assert 0 < len(list(wb["disputes"].iter_rows(min_row=2))) <= 60
    head = [c.value for c in wb["audit"][1]]
    for r in audit:
        for f in CORE:
            assert r[head.index(f"{f}_final")] == r[head.index(f"{f}_prefill")]


def test_final_columns_have_dropdowns():
    wb = load_workbook(BOOK)
    ws = wb["audit"]
    head = [c.value for c in ws[1]]
    covered = {}
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            covered[rng.min_col] = dv
    for f in CORE:
        j = head.index(f"{f}_final") + 1
        if f in FREE:
            assert j not in covered, f
        else:
            dv = covered[j]
            assert dv.type == "list" and dv.showErrorMessage and dv.errorStyle == "stop", f
    # prefilled values are always inside the allowed list
    for r in ws.iter_rows(min_row=2, values_only=True):
        for f in CORE:
            if f in FREE:
                continue
            dv = covered[head.index(f"{f}_final") + 1]
            allowed = dv.formula1.strip('"').split(",")
            assert str(r[head.index(f"{f}_prefill")]) in allowed, (f, r[head.index(f"{f}_prefill")])


def test_viewer_offline_and_blind():
    s = (ANN / "viewer_audit.html").read_text(encoding="utf-8")
    assert '"id": "u001"' in s and not re.search(r"(?:src|href)=\"https?://|<link ", s)
    assert not re.search(r"llm_AB|\"rule\"|规则|模型", s)
