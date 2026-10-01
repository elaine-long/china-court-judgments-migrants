import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src import paths
from src.annotation_schema import COLUMNS, TAIL

ANN = paths.ROOT / "annotation"
BOOKS = [(ANN / "T05_annotation.xlsx", "ann_id", ["main", "pilot"]), (ANN / "T05_recheck.xlsx", "rid", ["main"])]
pytestmark = pytest.mark.skipif(not (ANN / "T05_annotation.xlsx").exists(), reason="workbooks not built")

# rule/LLM output column names that must never appear in an annotation sheet
FORBIDDEN = {"doc_id", "stratum", "weight", "migrant_city", "migrant_city_nores", "d1_term_months", "amt_yuan",
             "pretrial_status_at_judgment", "leniency_proc", "pilot_city"}


def _allowed(wb, dv):
    f = dv.formula1
    if f.startswith('"'):
        return f.strip('"').split(",")
    name = f.lstrip("=")
    dn = wb.defined_names[name]
    ((sheet, rng),) = list(dn.destinations)
    return [c[0].value for c in wb[sheet][rng.replace("$", "")]]


@pytest.mark.parametrize("path,id_col,sheets", BOOKS)
def test_structure_and_validation(path, id_col, sheets):
    wb = load_workbook(path)
    for name in sheets:
        ws = wb[name]
        head = [c.value for c in ws[1]]
        assert head[:3] == [id_col, "court_name", "judgment_date"]
        assert head[3:] == [c[0] for c in COLUMNS + TAIL]
        assert not FORBIDDEN & set(head)
        assert ws.freeze_panes == "D2"
        n = ws.max_row
        by_col = {}
        for dv in ws.data_validations.dataValidation:
            for rng in dv.sqref.ranges:
                by_col[rng.min_col] = (dv, rng)
        for j, (col, allowed, _, _) in enumerate(COLUMNS + TAIL, start=4):
            if allowed:
                dv, rng = by_col[j]
                assert dv.type == "list" and dv.showErrorMessage and dv.errorStyle == "stop", col
                assert rng.min_row == 2 and rng.max_row >= n, col
                assert _allowed(wb, dv) == allowed, col
            else:
                assert j not in by_col, col
        assert "denial_reason_cat" in head
        cm = ws.cell(row=1, column=head.index("denial_reason_cat") + 1).comment
        assert cm and all(code in cm.text for code in ["residence", "record", "severity", "attitude", "generic"])
        # annotation cells start empty
        for row in ws.iter_rows(min_row=2, min_col=4, values_only=True):
            assert all(v is None for v in row)
    assert {"codebook", "log"} <= set(wb.sheetnames)


def test_row_counts_and_ids():
    import pandas as pd
    wb = load_workbook(ANN / "T05_annotation.xlsx")
    ids = [r[0] for r in wb["main"].iter_rows(min_row=2, values_only=True)]
    assert ids == sorted(pd.read_csv(ANN / "sample_main.csv").ann_id) and len(ids) == 300
    assert len(list(wb["pilot"].iter_rows(min_row=2))) == 10
    rc = load_workbook(ANN / "T05_recheck.xlsx")["main"]
    rids = [r[0] for r in rc.iter_rows(min_row=2, values_only=True)]
    assert rids == [f"rid_{i:03d}" for i in range(1, 61)]


def test_viewers_offline_and_blind():
    import re
    for name, key in [("viewer_main.html", "a001"), ("viewer_pilot.html", "p001"), ("viewer_recheck.html", "rid_001")]:
        s = (ANN / name).read_text(encoding="utf-8")
        assert key in s
        assert not re.search(r"(?:src|href)=\"https?://|@import|<link ", s)  # no external resources
        assert "pilot_city" not in s and "stratum" not in s and "P2_plea" not in s
    rc = (ANN / "viewer_recheck.html").read_text(encoding="utf-8")
    assert not re.search(r'"id": "a\d{3}"', rc)  # recheck viewer shows rid only
