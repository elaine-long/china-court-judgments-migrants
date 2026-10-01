"""T04 steps 2-3: annotation workbooks (openpyxl) and offline HTML viewers.

No rule or LLM extraction results are written to the workbooks or viewers.
"""
import html
import json
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src import paths  # noqa: E402
from src.annotation_schema import COLUMNS, DENIAL_CODES, TAIL  # noqa: E402
from src.extract_rules import normalize  # noqa: E402
from src.io import load_meta, load_texts  # noqa: E402
from src.segment import find_spans  # noqa: E402

ANN = paths.ROOT / "annotation"
ALL_COLS = COLUMNS + TAIL
HEAD_FILL = PatternFill("solid", fgColor="DDE7F0")


def doc_info(doc_ids):
    meta = load_meta(["doc_id", "court_name", "case_number"]).set_index("doc_id")
    txt = load_texts(doc_ids, columns=["judgment", "judgment_date"]).set_index("doc_id")
    out = txt.join(meta, how="left")
    return out


# ----------------------------------------------------------------------------- workbook
def _annotation_sheet(ws, id_col, ids, info):
    head = [id_col, "court_name", "judgment_date"] + [c[0] for c in ALL_COLS]
    ws.append(head)
    for i, (key, did) in enumerate(ids):
        r = info.loc[did]
        ws.append([key, r.court_name, r.judgment_date] + [None] * len(ALL_COLS))
    n = len(ids) + 1
    for j, h in enumerate(head, start=1):
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    lists = {}
    for j, (col, allowed, vals, rule) in enumerate(ALL_COLS, start=4):
        letter = get_column_letter(j)
        text = f"{vals}\n{rule}"
        if col == "denial_reason_cat":
            text += "\n可多选，分号分隔。代码：" + "; ".join(DENIAL_CODES) + "；无则 NA"
        ws.cell(row=1, column=j).comment = Comment(text, "codebook v1", width=300, height=140)
        if allowed:
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True,
                                showErrorMessage=True, errorStyle="stop",
                                errorTitle="不在允许取值内", error="允许取值：" + " / ".join(allowed))
            if len(dv.formula1) > 255:  # long lists go to the hidden `lists` sheet
                lists[col] = allowed
                dv.formula1 = f"=lists_{col}"
            dv.add(f"{letter}2:{letter}{max(n, 2)}")
            ws.add_data_validation(dv)
        width = 16 if (allowed or col in ("minutes", "flag")) else 24
        ws.column_dimensions[letter].width = width
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "D2"
    return lists


def build_workbook(path, id_col, ids, info, pilot=None, extra_sheets=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "main"
    lists = _annotation_sheet(ws, id_col, ids, info)
    if pilot is not None:
        lists.update(_annotation_sheet(wb.create_sheet("pilot"), id_col, pilot, info))
    if extra_sheets:
        cb = wb.create_sheet("codebook")
        cb.append(["列名", "取值", "一句话规则"])
        for col, _, vals, rule in ALL_COLS:
            cb.append([col, vals, rule])
        for c, w in zip("ABC", (22, 34, 90)):
            cb.column_dimensions[c].width = w
        for cell in cb[1]:
            cell.font = Font(bold=True)
        cb.freeze_panes = "A2"
        lg = wb.create_sheet("log")
        lg.append(["batch", "date", "n_docs", "minutes_total", "n_flag"])
        for b in range(1, 11):
            lg.append([b, None, None, None, None])
        for cell in lg[1]:
            cell.font = Font(bold=True)
        lg.freeze_panes = "A2"
    if lists:
        from openpyxl.workbook.defined_name import DefinedName
        ls = wb.create_sheet("lists")
        for j, (col, vals) in enumerate(lists.items(), start=1):
            for i, v in enumerate(vals, start=1):
                ls.cell(row=i, column=j, value=v)
            letter = get_column_letter(j)
            wb.defined_names[f"lists_{col}"] = DefinedName(
                f"lists_{col}", attr_text=f"lists!${letter}$1:${letter}${len(vals)}")
        ls.sheet_state = "hidden"
    wb.save(path)


# ----------------------------------------------------------------------------- viewer
SEG_TITLES = {"header": "当事人", "charge": "指控", "facts": "查明", "reasoning": "本院认为", "verdict": "判决"}
HL = re.compile(r"户籍|户口|出生|生于|住|(?<=[省市县区州旗])人(?=[，,。；;])|拘留|逮捕|取保|监视居住|羁押|缓刑|"
                r"认罪认罚|具结|速裁|辩护人|值班律师")


def _hl(s):
    return HL.sub(lambda m: f"<mark>{m.group(0)}</mark>", html.escape(s))


def blocks(text):
    t = normalize(text)
    sp = find_spans(t)
    segs = sorted((v[0], v[1], k) for k, v in sp.items() if v)
    if not segs:
        return [("全文（分段失败）", _hl(t))]
    out, pos = [], 0
    for s, e, k in segs:
        if s > pos:
            out.append(("其他", _hl(t[pos:s])))
        s = max(s, pos)
        if e > s:
            out.append((SEG_TITLES[k], _hl(t[s:e])))
        pos = max(pos, e)
    if pos < len(t):
        out.append(("其他", _hl(t[pos:])))
    return out


VIEWER = """<!doctype html>
<html lang="zh"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title>
<style>
:root{--bg:#fbfbf9;--fg:#1f2328;--muted:#6b7280;--line:#e5e7eb;--side:#f3f4f6;--mark:#fff1b8;--sel:#dbeafe}
@media (prefers-color-scheme:dark){:root{--bg:#16181c;--fg:#e6e6e6;--muted:#9aa0a6;--line:#2c3036;--side:#1d2025;--mark:#5a4a12;--sel:#1e3a5f}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.75 "PingFang SC","Hiragino Sans GB","Microsoft YaHei",sans-serif;display:flex;height:100vh}
nav{width:150px;flex:none;overflow-y:auto;background:var(--side);border-right:1px solid var(--line)}
nav input{width:calc(100% - 16px);margin:8px;padding:4px 6px;border:1px solid var(--line);border-radius:4px;background:var(--bg);color:var(--fg)}
nav a{display:block;padding:4px 12px;color:var(--fg);text-decoration:none;font-family:ui-monospace,Menlo,monospace;font-size:13px;cursor:pointer}
nav a.on{background:var(--sel)}main{flex:1;overflow-y:auto;padding:20px 32px 60px;max-width:980px}
.top{position:sticky;top:0;background:var(--bg);padding:8px 0 10px;border-bottom:1px solid var(--line);margin-bottom:12px}
.top .id{font:600 18px ui-monospace,Menlo,monospace}.top .meta{color:var(--muted);font-size:14px}
h3{font-size:13px;letter-spacing:.08em;color:var(--muted);margin:22px 0 4px;border-bottom:1px solid var(--line);padding-bottom:2px}
p.seg{margin:0;white-space:pre-wrap;word-break:break-all}mark{background:var(--mark);color:inherit;border-radius:2px;padding:0 1px}
.hint{color:var(--muted);font-size:12px;margin-top:4px}
@media (max-width:700px){body{flex-direction:column}nav{width:100%;height:30vh}main{padding:12px 16px}}
</style></head><body>
<nav><input id="q" placeholder="跳转 __IDNAME__"><div id="list"></div></nav>
<main><div class="top"><div class="id" id="id"></div><div class="meta" id="meta"></div>
<div class="hint">浅色底只是关键词提示，不代表任何判断。键盘 ↑/↓ 切换。</div></div><div id="body"></div></main>
<script>
const DOCS=__DATA__;
const list=document.getElementById('list');let cur=0;
DOCS.forEach((d,i)=>{const a=document.createElement('a');a.textContent=d.id;a.onclick=()=>show(i);list.appendChild(a);});
function show(i){cur=i;const d=DOCS[i];
 document.getElementById('id').textContent=d.id;
 document.getElementById('meta').textContent=d.court+'　'+d.date+'　'+d.cn;
 document.getElementById('body').innerHTML=d.blocks.map(b=>'<h3>'+b[0]+'</h3><p class="seg">'+b[1]+'</p>').join('');
 [...list.children].forEach((a,j)=>a.classList.toggle('on',j===i));
 list.children[i].scrollIntoView({block:'nearest'});document.querySelector('main').scrollTop=0;
 history.replaceState(null,'','#'+d.id);}
document.getElementById('q').addEventListener('keydown',e=>{if(e.key==='Enter'){const k=DOCS.findIndex(d=>d.id===e.target.value.trim());if(k>=0)show(k);}});
document.addEventListener('keydown',e=>{if(e.target.tagName==='INPUT')return;
 if(e.key==='ArrowDown'&&cur<DOCS.length-1){show(cur+1);e.preventDefault();}
 if(e.key==='ArrowUp'&&cur>0){show(cur-1);e.preventDefault();}});
const h=location.hash.slice(1);const k=DOCS.findIndex(d=>d.id===h);show(k>=0?k:0);
</script></body></html>"""


def build_viewer(path, title, idname, ids, info):
    docs = []
    for key, did in ids:
        r = info.loc[did]
        docs.append({"id": key, "court": r.court_name or "", "date": r.judgment_date or "",
                     "cn": r.case_number or "", "blocks": blocks(r.judgment or "")})
    data = json.dumps(docs, ensure_ascii=False).replace("</", "<\\/")
    page = VIEWER.replace("__TITLE__", title).replace("__IDNAME__", idname).replace("__DATA__", data)
    path.write_text(page, encoding="utf-8")


def main():
    main_s = pd.read_csv(ANN / "sample_main.csv").sort_values("ann_id")
    pilot_s = pd.read_csv(ANN / "sample_pilot.csv").sort_values("ann_id")
    remap = pd.read_csv(ANN / "_private" / "remap.csv").sort_values("rid")
    info = doc_info(pd.concat([main_s.doc_id, pilot_s.doc_id]).unique())
    m_ids = list(zip(main_s.ann_id, main_s.doc_id))
    p_ids = list(zip(pilot_s.ann_id, pilot_s.doc_id))
    r_ids = list(zip(remap.rid, remap.doc_id))
    build_workbook(ANN / "T05_annotation.xlsx", "ann_id", m_ids, info, pilot=p_ids)
    build_workbook(ANN / "T05_recheck.xlsx", "rid", r_ids, info, extra_sheets=True)
    build_viewer(ANN / "viewer_main.html", "标注阅读器 · main", "ann_id", m_ids, info)
    build_viewer(ANN / "viewer_pilot.html", "标注阅读器 · pilot", "ann_id", p_ids, info)
    build_viewer(ANN / "viewer_recheck.html", "标注阅读器 · recheck", "rid", r_ids, info)
    for f in sorted(ANN.glob("*")):
        if f.is_file():
            print(f.name, f.stat().st_size)


if __name__ == "__main__":
    main()
