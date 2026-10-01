"""T06b §4-5: hybrid analysis table and the T05 audit materials.

    python scripts/T06b_post.py hybrid   # data/analysis_hybrid.parquet (+ refreshed analysis_rules)
    python scripts/T06b_post.py audit    # annotation/T05_audit.xlsx, viewer_audit.html, README.md

Hybrid rule (T06b §4): rule value if present; else the R fill when models A and B agree; else missing.
Every filled variable gets a *_source column: rule / llm_AB / NA.
"""
import html
import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from src import paths  # noqa: E402
from src.extract_rules import normalize  # noqa: E402
from src.io import load_meta, load_texts  # noqa: E402
from src.segment import find_spans  # noqa: E402
from T09a_build_panel import build, load_all  # noqa: E402

T = paths.TABLES
ANN = paths.ROOT / "annotation"


# ----------------------------------------------------------------------------- §4 hybrid
def _src(rule_col, filled_mask):
    return np.where(rule_col.notna(), "rule", np.where(filled_mask, "llm_AB", "NA"))


def hybrid():
    df = load_all()
    d = build(df)                                     # all theft/universe/刑初/2013-2019 rows, no filters
    rf = pd.read_parquet(paths.DATA / "R_fill.parquet")
    fill = rf[rf.filled_by == "A=B"].pivot_table(index="doc_id", columns="field", values="filled", aggfunc="first")
    fill = fill.reindex(d.doc_id)
    fill.index = d.index

    h = d.copy()
    # migrant (main = migrant_city_nores): R fills origin via hukou/birthplace/native text (prompt v2 never
    # uses residence for origin_text), so the filled value is on the migrant_city_nores definition
    f_mig = pd.to_numeric(fill.get("nonlocal_origin"), errors="coerce").astype(float)
    use = h.migrant.isna() & f_mig.notna()
    h.loc[use, "migrant"] = f_mig[use]
    h.loc[use, "migrant_city_nores"] = f_mig[use]
    h["migrant_source"] = _src(d.migrant, use)
    f_res = pd.to_numeric(fill.get("residence_local"), errors="coerce").astype(float)
    use_r = h.local_residence.isna() & f_res.notna()
    h.loc[use_r, "local_residence"] = f_res[use_r]
    h["local_residence_source"] = _src(d.local_residence, use_r)
    # penalty fields: only where the rule outcome was incomplete
    pen = fill.get("penalty_type")
    term = pd.to_numeric(fill.get("term_months"), errors="coerce").astype(float)
    prob = pd.to_numeric(fill.get("probation"), errors="coerce").astype(float)
    need_pen = ~d.outcomes_ok
    pen_ok = pen.notna() & (~pen.isin(["有期徒刑", "拘役"]) | term.notna()) if pen is not None else pd.Series(False, index=d.index)
    use_p = need_pen & pen_ok
    h.loc[use_p, "term_months"] = term[use_p].where(pen[use_p].isin(["有期徒刑", "拘役"]))
    h.loc[use_p, "log_term"] = np.log1p(h.loc[use_p, "term_months"])
    h.loc[use_p & h.probation.isna() & prob.notna(), "probation"] = prob[use_p & h.probation.isna() & prob.notna()]
    h["outcomes_ok"] = d.outcomes_ok | use_p
    h["outcomes_source"] = np.where(d.outcomes_ok, "rule", np.where(use_p, "llm_AB", "NA"))

    def pap(x):
        return x[x.outcomes_ok & x.city.notna() & x.migrant.notna()]

    rules = pap(d)
    hyb = pap(h)
    rules.to_parquet(paths.DATA / "analysis_rules.parquet", index=False)      # refreshed with T06b rules
    # unfiltered-by-migrant versions (for coverage, IPW and alternative migrant definitions; T09b)
    h[h.outcomes_ok & h.city.notna()].to_parquet(paths.DATA / "analysis_hybrid_full.parquet", index=False)
    d[d.outcomes_ok & d.city.notna()].to_parquet(paths.DATA / "analysis_rules_full.parquet", index=False)
    hyb.to_parquet(paths.DATA / "analysis_hybrid.parquet", index=False)
    base = d[d.city.notna()]
    hb = h[h.city.notna()]
    cov = pd.DataFrame([
        {"variable": "migrant (migrant_city_nores), defined", "rules": base.migrant.notna().mean(),
         "hybrid": hb.migrant.notna().mean(), "n_filled": int(use.sum())},
        {"variable": "local_residence, defined", "rules": base.local_residence.notna().mean(),
         "hybrid": hb.local_residence.notna().mean(), "n_filled": int(use_r.sum())},
        {"variable": "sentence fields complete (outcomes_ok)", "rules": base.outcomes_ok.mean(),
         "hybrid": hb.outcomes_ok.mean(), "n_filled": int(use_p.sum())},
        {"variable": "PAP main sample (rows)", "rules": len(rules), "hybrid": len(hyb),
         "n_filled": len(hyb) - len(rules)},
    ])
    cov["change"] = cov.hybrid - cov.rules
    cov.to_csv(T / "T06b_hybrid_coverage.csv", index=False)
    print(cov.to_string(index=False))
    print(hyb.migrant_source.value_counts().to_string())


# ----------------------------------------------------------------------------- §5 audit
CORE = [  # codebook v2 core fields: (name, allowed values for the final-value dropdown or None)
    ("origin_text", None), ("nonlocal_origin", ["1", "0", "NA"]), ("residence_local", ["1", "0", "NA"]),
    ("status_at_judgment", ["在押", "取保", "监视居住", "其他", "NA"]), ("ever_bail", ["1", "0"]),
    ("plea_formal", ["1", "0"]), ("jiejie", ["1", "0"]), ("procedure", ["速裁", "简易", "普通", "NA"]),
    ("counsel_any", ["1", "0"]),
    ("penalty_type", ["有期徒刑", "拘役", "管制", "单处罚金", "免予刑事处罚", "其他"]),
    ("term_months", None), ("probation", ["1", "0"]),
]
_CUES = {"status_at_judgment": r"现[^。；]{0,20}(?:羁押|押|取保|监视居住|在家|关押)|羁押于[^。；]{0,15}"
                               r"|[^，。；]{0,12}被(?:监视居住|取保候审|执行逮捕|逮捕|刑事拘留|羁押)",
         "ever_bail": r"取保候审", "plea_formal": r"认罪认罚", "jiejie": r"具结",
         "procedure": r"适用[^。；]{0,8}程序", "counsel_any": r"辩护人[^，。；]{0,12}",
         "penalty_type": r"判处[^。；]{0,20}", "term_months": r"判处[^。；]{0,20}", "probation": r"缓刑[^。；]{0,12}",
         "origin_text": r"(?:户籍|户口|出生于|生于)[^，。；]{0,20}", "nonlocal_origin": r"(?:户籍|户口|出生于|生于)[^，。；]{0,20}",
         "residence_local": r"住[^，。；]{0,20}"}


_VERDICT_FIELDS = {"penalty_type", "term_months", "probation"}
_PARTY_FIELDS = {"origin_text", "nonlocal_origin", "residence_local", "status_at_judgment", "ever_bail", "counsel_any"}
_ORIGIN_CUE = re.compile(r"(?:户籍|户口|出生于|生于|出生地)[^，。；]{0,6}$")


def _evidence(text, field, value):
    """Short neutral excerpt (≤ 30 chars) supporting the prefilled value; same method for every source.

    Section-aware: sentence fields are looked up in the verdict only (not in prior convictions), party
    fields in the party + procedural block (not in the court name or the evidence list)."""
    t = normalize(text)
    sp = find_spans(t)
    nxt = [v[0] for k, v in sp.items() if v and k != "header"]
    party = t[:min(nxt)] if nxt else t[:1500]
    region = (t[sp["verdict"][0]:sp["verdict"][1]] if sp.get("verdict") else t) if field in _VERDICT_FIELDS \
        else (party if field in _PARTY_FIELDS else t)
    if field in ("origin_text", "nonlocal_origin") and isinstance(value, str) and value not in ("NA", "0", "1"):
        for m in re.finditer(re.escape(value), region):
            if _ORIGIN_CUE.search(region[max(0, m.start() - 12):m.start()]) or \
                    region[m.end():m.end() + 1] == "人":
                return region[max(0, m.start() - 6):m.end() + 2][:30]
    m = re.search(_CUES.get(field, r"$^"), region)
    return m.group(0)[:30] if m else ""


def prefill_table(doc_ids):
    """Hybrid values for the 12 core fields: rule first; word fields rule only; if the rule is NA and the
    two models agree on V, the model value (T06b §4 rule applied to V)."""
    from T06b_run import rule_values  # noqa: E402
    import duckdb
    R = rule_values(doc_ids)
    con = duckdb.connect()
    con.execute("SET enable_progress_bar = false")
    con.register("ids", pd.DataFrame({"doc_id": list(doc_ids)}))
    w = con.execute(f"""SELECT doc_id, plea_formal, jiejie, procedure, d1_counsel FROM read_parquet('{paths.DATA / "feat_rules.parquet"}')
                        WHERE doc_id IN (SELECT doc_id FROM ids)""").df().set_index("doc_id")
    vl = pd.read_parquet(paths.DATA / "V_labels.parquet")
    vl = vl[vl.doc_id.isin(doc_ids)].set_index(["doc_id", "field"])
    out, src = {}, {}
    for d in doc_ids:
        row, srow = {}, {}
        for f, _ in CORE:
            if f in ("plea_formal", "jiejie"):
                val, s = ("1" if w.at[d, f] else "0"), "rule"
            elif f == "procedure":
                val, s = (w.at[d, "procedure"] if isinstance(w.at[d, "procedure"], str) else "NA"), "rule"
            elif f == "counsel_any":
                val, s = ("1" if w.at[d, "d1_counsel"] else "0"), "rule"
            else:
                rv = R.at[d, f] if f in R.columns else "NA"
                val, s = rv, "rule"
                if (rv in (None, "NA") or pd.isna(rv)) and (d, f) in vl.index:
                    a, b = vl.at[(d, f), "A"], vl.at[(d, f), "B"]
                    if pd.notna(a) and pd.notna(b) and a == b and a != "NA":
                        val, s = a, "llm_AB"
                if val in (None,) or (isinstance(val, float) and np.isnan(val)):
                    val = "NA"
            if f == "term_months" and val not in ("NA", None):
                val = f"{float(val):g}"
            row[f], srow[f] = val, s
        out[d], src[d] = row, srow
    return out, src


def audit():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    v = pd.read_parquet(paths.DATA / "V_sample.parquet")
    rng = np.random.default_rng(paths.SEED)
    n_total = 150
    alloc = {s: n_total * len(g) // len(v) for s, g in v.groupby("stratum")}
    for s in rng.choice(sorted(alloc), n_total - sum(alloc.values()), replace=False):
        alloc[s] += 1
    picks = pd.concat([g.sample(alloc[s], random_state=paths.SEED) for s, g in v.groupby("stratum")])
    picks = picks.sample(frac=1, random_state=paths.SEED).reset_index(drop=True)
    picks.insert(0, "uid", [f"u{i:03d}" for i in range(1, len(picks) + 1)])
    picks["n_audit_stratum"] = picks.stratum.map(alloc)
    picks["p_audit_given_V"] = picks.n_audit_stratum / 625
    picks["p_audit_total"] = picks.n_audit_stratum / picks.N_stratum
    (ANN / "_private").mkdir(parents=True, exist_ok=True)
    picks.to_csv(ANN / "_private" / "audit_sample.csv", index=False)
    picks[["uid", "stratum", "p_audit_given_V", "p_audit_total"]].to_csv(ANN / "sample_audit.csv", index=False)

    # disputes: V items where B's independent value differs from the post-adjudication value, or low confidence
    vl = pd.read_parquet(paths.DATA / "V_labels.parquet")
    adj = vl[vl.B_adj_side.notna()] if "B_adj_side" in vl else vl.iloc[0:0]
    adj = adj.assign(post=np.where(adj.B_adj_side == "rule", adj.rule,
                                   np.where(adj.B_adj_side == "A", adj.A, adj.B_adj_value)))
    disp = adj[(adj.post.astype(str) != adj.B.astype(str)) | (adj.B_adj_conf == "低")]
    disp = disp[~disp.doc_id.isin(picks.doc_id)]
    disp_docs = (disp.groupby("doc_id").field.apply(lambda s: ";".join(sorted(set(s)))).reset_index()
                 .sample(frac=1, random_state=paths.SEED).head(60).reset_index(drop=True))
    disp_docs.insert(0, "uid", [f"d{i:03d}" for i in range(1, len(disp_docs) + 1)])
    disp_docs.to_csv(ANN / "_private" / "disputes.csv", index=False)
    print(f"audit {len(picks)} docs; disputes candidates {disp.doc_id.nunique()} docs, taken {len(disp_docs)}")

    all_ids = list(picks.doc_id) + list(disp_docs.doc_id)
    pre, src = prefill_table(all_ids)
    meta = load_meta(["doc_id", "court_name", "case_number"]).set_index("doc_id")
    txt = load_texts(all_ids, columns=["judgment", "judgment_date"]).set_index("doc_id")
    pd.DataFrame([{"doc_id": d, **{f"src_{f}": s for f, s in src[d].items()}} for d in all_ids]) \
        .to_csv(ANN / "_private" / "audit_prefill_sources.csv", index=False)

    head = ["uid", "court_name", "judgment_date"]
    for f, _ in CORE:
        head += [f"{f}_prefill", f"{f}_final", f"{f}_evidence"]
    head += ["changed", "flag", "note", "minutes"]
    evid = {}

    def sheet(ws, rows, focus=None):
        ws.append(head)
        for uid, d in rows:
            r = [uid, meta.at[d, "court_name"], txt.at[d, "judgment_date"]]
            for f, _ in CORE:
                val = pre[d][f]
                e = _evidence(txt.at[d, "judgment"], f, val)
                evid.setdefault(uid, []).append(e)
                r += [val, val, e]
            r += [0, None, (f"待核字段：{focus[d]}" if focus else None), None]
            ws.append(r)
        n = len(rows) + 1
        for j, h in enumerate(head, start=1):
            c = ws.cell(row=1, column=j)
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if h.endswith("_final"):
                c.fill = PatternFill("solid", fgColor="E2EFDA")
            ws.column_dimensions[get_column_letter(j)].width = 22 if h.endswith("_evidence") else 12
        for j, (f, allowed) in enumerate(CORE):
            col = get_column_letter(4 + 3 * j + 1)        # the *_final column
            if allowed:
                dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=False,
                                    showErrorMessage=True, errorStyle="stop", error="只允许：" + " / ".join(allowed))
                dv.add(f"{col}2:{col}{max(n, 2)}")
                ws.add_data_validation(dv)
        for colname, allowed in (("changed", ["1", "0"]), ("flag", ["1"])):
            col = get_column_letter(head.index(colname) + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
            dv.add(f"{col}2:{col}{max(n, 2)}")
            ws.add_data_validation(dv)
        ws.freeze_panes = "D2"
        ws.cell(row=1, column=head.index("changed") + 1).comment = Comment(
            "任何一个 final 值改动了就填 1", "T06b")

    wb = Workbook()
    ws = wb.active
    ws.title = "disputes"
    sheet(ws, list(zip(disp_docs.uid, disp_docs.doc_id)), focus=dict(zip(disp_docs.doc_id, disp_docs.field)))
    sheet(wb.create_sheet("audit"), list(zip(picks.uid, picks.doc_id)))
    cb = wb.create_sheet("codebook")
    cb.append(["字段", "取值", "规则（手册 v2）"])
    rules_txt = {
        "origin_text": "户籍 > 出生地 > 籍贯中第一个非空的地名原文；不用住址", "nonlocal_origin": "来源地与法院不在同一地级市为 1",
        "residence_local": "住址在法院所在地级市为 1，外地为 0，没写为 NA",
        "status_at_judgment": "以“现……”句为准；没有时取最后一个事件；不能由刑期折抵套话推断",
        "ever_bail": "任何阶段取保过为 1", "plea_formal": "出现“认罪认罚”为 1", "jiejie": "出现“具结”为 1",
        "procedure": "以“适用……程序”明文为准；没写为 NA", "counsel_any": "第一被告人有辩护人为 1",
        "penalty_type": "以判决主文为准", "term_months": "主刑刑期（月），不含缓刑考验期", "probation": "宣告缓刑为 1"}
    for f, allowed in CORE:
        cb.append([f, " / ".join(allowed) if allowed else ("原文" if f == "origin_text" else "数字或 NA"), rules_txt[f]])
    for c, wdt in zip("ABC", (20, 40, 70)):
        cb.column_dimensions[c].width = wdt
    lg = wb.create_sheet("log")
    lg.append(["sheet", "batch", "date", "n_docs", "minutes_total", "n_changed", "n_flag"])
    wb.save(ANN / "T05_audit.xlsx")

    # viewer: evidence excerpts highlighted, no source shown
    build_viewer([*zip(disp_docs.uid, disp_docs.doc_id), *zip(picks.uid, picks.doc_id)], txt, meta, evid)
    (ANN / "README.md").write_text(README, encoding="utf-8")
    print("wrote annotation/T05_audit.xlsx, viewer_audit.html, README.md")


def build_viewer(ids, txt, meta, evid):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from T04_build_annotation import SEG_TITLES, VIEWER
    docs = []
    for uid, d in ids:
        t = normalize(txt.at[d, "judgment"])
        marks = sorted({e for e in evid.get(uid, []) if e and len(e) >= 2}, key=len, reverse=True)
        sp = find_spans(t)
        segs = sorted((v[0], v[1], k) for k, v in sp.items() if v)
        blocks, pos = [], 0
        for s, e, k in segs or [(0, len(t), None)]:
            if s > pos:
                blocks.append(("其他", t[pos:s]))
            blocks.append((SEG_TITLES.get(k, "全文"), t[max(s, pos):e]))
            pos = max(pos, e)
        if pos < len(t):
            blocks.append(("其他", t[pos:]))
        out = []
        for title, b in blocks:
            h = html.escape(b)
            for m in marks:
                h = h.replace(html.escape(m), f"<mark>{html.escape(m)}</mark>")
            out.append((title, h))
        docs.append({"id": uid, "court": meta.at[d, "court_name"] or "", "date": txt.at[d, "judgment_date"] or "",
                     "cn": meta.at[d, "case_number"] or "", "blocks": out})
    data = json.dumps(docs, ensure_ascii=False).replace("</", "<\\/")
    page = (VIEWER.replace("__TITLE__", "审计阅读器 · T05").replace("__IDNAME__", "编号")
            .replace("浅色底只是关键词提示，不代表任何判断。", "浅色底是预填依据的原文位置，请回到全文核对。")
            .replace("__DATA__", data))
    (ANN / "viewer_audit.html").write_text(page, encoding="utf-8")


README = """# T05 人工审计：操作说明

**材料**：`T05_audit.xlsx`（`disputes`、`audit`、`codebook`、`log` 四张表）＋ `viewer_audit.html`（离线阅读器，左侧按编号跳转）。

## 顺序
1. **先做 `disputes`**（d001 起，最多 60 份）：两个模型意见不一致或置信度低的文书。`note` 列写了"待核字段"，先核这些字段，其余字段顺手核对。
2. **再做 `audit`**（u001–u150）：随机审计样本，是准确率和 DSL 纠偏的依据，**每个字段都要核对**。

## 每份怎么做
- 每个字段三列：`_prefill`（预填）、`_final`（最终值，默认等于预填）、`_evidence`（原文摘录）。
- 在阅读器里打开同一编号，**回到原文核对**，不能只看摘录就确认。浅色底标出的是摘录在原文中的位置。
- 预填正确：不动。预填错误或缺失：改 `_final`（有下拉的只能从下拉选），并把 `changed` 填 1。
- 拿不准：`flag` 填 1，`note` 写一句原因，继续往下。
- 每份在 `minutes` 记用时；每批结束在 `log` 表记一行。
- 字段口径见 `codebook` 表与 `docs/codebook.md` v2。

## 注意
- 表里不显示预填值来自规则还是模型，请不要去 `data/` 或 `annotation/_private/` 查看。
- 预计用时：disputes 约 1 小时，audit 约 2.5–3 小时，合计约 3–4 小时。
"""


if __name__ == "__main__":
    {"hybrid": hybrid, "audit": audit}[sys.argv[1]]()
